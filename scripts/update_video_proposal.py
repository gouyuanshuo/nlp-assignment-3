"""Update the existing course proposal; execute using the document artifact runtime."""

import io
import shutil
from pathlib import Path

from docx import Document
from docx.shared import RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
PROPOSAL = Path(
    "/home/gys/Course/2026-Spring-42173-Advanced-Natural-Language-Processing/04_Assignments/Assignment_03_Cross_Modal_Conflict/Plan/Assignment3_Transformer_Fusion_CMU_MOSEI_Project_Plan.docx"
)


def main():
    backup = ROOT / "data/proposal-backup/before-video-prototype.docx"
    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(PROPOSAL, backup)
    doc = Document(backup)
    replacements = {
        0: "Transformer Fusion for English Video Sentiment Classification Project Plan",
        2: "We will develop a prototype that accepts a newly uploaded English video and classifies its overall sentiment. Frozen text, acoustic and facial-feature extractors will produce sequences, which are aligned to speech words, projected into a common embedding dimension and combined by a Transformer. CMU-MOSEI will be the sole research dataset. We will compare this system with text-only prediction and pooled feature concatenation, measure sensitivity to missing or corrupted features, and demonstrate the complete video-to-prediction workflow in a local browser interface.",
        12: "Our contribution is a reproducible application and controlled evaluation of a compact Transformer fusion classifier, together with a working video-upload prototype. We will test whether retaining sequences improves on pooled concatenation and identify which modalities support prediction under clean and perturbed conditions. Transformer fusion is an established method; our claim concerns the implemented system and its evaluation. Benchmark scores on released features and the behavior of the live speech-recognition pipeline will be reported separately.",
        17: "Text. Regenerate 768-dimensional word features from the released timestamped English words with a pinned, frozen BERT-base-uncased model. Mean-pool the final hidden states of each word's wordpieces. For new videos, obtain word timestamps with local faster-whisper small.en, then use the same BERT encoder and pooling. The text-only baseline uses the same features as both fusion models.",
        18: "Vision. Select 40 named OpenFace 2 features from the larger CSD release: two gaze angles, three head-rotation angles, 17 action-unit intensities and 18 action-unit presence indicators. Record names, order and units, and retain frames with successful tracking and confidence at least 0.8. The live extractor uses a pinned OpenFace image with the same feature families. Camera-dependent landmark coordinates and head translations are excluded.",
        19: "Audio. Select 73 named COVAREP features from the same CSD release. The live implementation runs the pinned upstream algorithms in Octave with tested compatibility adapters. Exclude the MATLAB-toolbox-dependent creak channel consistently from training and inference. Record 16 kHz mono input, 10 ms feature intervals, native feature units and adapter versions. A feature manifest fixes the 768/73/40 input widths and exact column order.",
        20: "Keep the feature encoders frozen and train the projection, fusion and classification layers. Align finite acoustic and visual measurements to word intervals by overlap-weighted averaging; quality failures receive explicit masks. Divide each source clip into windows of at most 50 words, also respecting BERT's 512-token limit. Retain every window. Fit normalization only on valid, non-neutral training tokens, and apply those saved statistics unchanged to validation, test and uploads.",
        22: "Apply a separate learned projection to each modality so that every feature vector has a common width, initially 256. Add learned modality embeddings and within-window positional information, then concatenate the three sequences along the token axis. Word alignment supplies the shared temporal reference. Each window preserves valid lengths and masks, including the distinction between missing measurements and padded positions.",
        26: "Use the final CLS representation as a fused window embedding. A small MLP outputs two logits for each window. Take the duration-weighted mean of those logits across all windows from the source clip, then apply softmax and argmax for one clip prediction. Masked mean pooling is a planned ablation. The same window aggregation applies to both baselines and uploaded videos.",
        27: "Train with mean cross-entropy on the aggregated clip logits and one binary label per source clip. A long clip therefore remains one supervised example, irrespective of how many windows it contains. Gradients pass through every window. Each baseline uses the same target, loss and aggregation; the core model has one fused classification head.",
        33: "Use the larger samarwarsi/cmu-mosei Kaggle package, version 1, selecting timestamped words, sentiment labels, COVAREP and OpenFace 2 CSD streams, totaling about 28.4 GB. This is a timestamped feature package; raw videos are not included. Preserve the official SDK video folds, parsed from a pinned source revision without executing downloaded code. Validate roots, column schemas, intervals, source-video membership and locally calculated file hashes. Report neutral exclusions, missing streams and clips without words. Keep independent label segments in their source order, including genuine temporal overlaps. Retain the earlier MMSA aligned archive only as separate preliminary evidence.",
        39: "Robustness: separately remove text, audio or vision, and add Gaussian noise with standard deviations 0.1, 0.3 and 0.5 in normalized feature units to one modality at a time. Leave padding unchanged. Reuse the same source clips and deterministic clip/window noise realizations across models. Aggregate window logits before scoring each clip once and measure macro F1 loss relative to its clean result. Describe this as feature-level corruption.",
        49: "Phase 1. Obtain the four pinned CSD streams and official video folds. Audit labels, columns, intervals and exclusions. Process a real spoken video with the local extractors, confirm names and units against the released streams, and freeze the shared alignment, windowing and feature manifest. Exact source extractor commits are unavailable, so numerical identity to the historical release will not be asserted.",
        50: "Phase 2. Regenerate BERT words and construct a memory-mapped cache with training-only normalization. Train the text-only and pooled-concatenation baselines using clip-level cross-entropy and validation-based checkpoint selection. Save the feature manifest, configuration and source hashes in each checkpoint. A small real-data integration subset checks execution before the complete research folds are used.",
        51: "Phase 3. Train the compact Transformer with shared projections, modality and position embeddings, joint attention and clip-level window aggregation. Verify padding invariance, finite gradients, complete window coverage, one evaluation record per clip and reloadable checkpoints. Use an available Google Colab GPU, preferably A100, for preparation and full training.",
        52: "Phase 4. Measure one Transformer epoch on the allocated Colab GPU, record the displayed compute-unit rate and estimate the full matrix against the available 200-unit budget. Freeze configurations and run three seeds for the core models and planned ablations. Save optimizer and random-number state at every completed epoch to Drive, allowing interrupted runs to resume. Evaluate the fixed checkpoints with matched feature corruption and source-video bootstrap comparisons.",
        53: "Phase 5. Export a compatible checkpoint, selected by validation results, to the local video-upload prototype. Verify successful uploads, invalid or silent clips, missing-face errors and temporary-file cleanup. Prepare the 2,500-3,500 word report in an 11-point font and the oral presentation, with traceable experiments and individual contributions. Report model-training time and complete upload-processing latency separately.",
        58: "Feature validation and robustness: verify COVAREP/OpenFace extraction, feature names and units, timestamp alignment, quality masks and window coverage; generate matched feature-noise and missing-modality conditions; document speech-recognition errors and train/upload distribution differences.",
        60: "Evaluation and integration: maintain the Colab notebooks and durable checkpoints, generate result tables and figures, inspect errors, build and validate the local upload interface, and integrate the report and presentation. Each member will record owned deliverables, code or experiment references and their individual contribution in the shared project tracker.",
        62: "Scope and compute. Use one English dataset and frozen feature encoders. Download on a CPU runtime where practical, cache prepared arrays, and use the actual Colab GPU and compute-unit rate to estimate cost. Availability and session duration can vary. Checkpoints in Drive preserve completed epochs. Complete the three core classifiers before optional comparators; any memory-related changes must be fixed consistently across comparisons.",
        66: "Leakage and preprocessing errors. Fit normalization on training data only, preserve official source-video partitions, audit duplicate IDs and masks, and retain one label per clip. Tests must cover word boundaries, missing measurements, late windows and checkpoint-manifest mismatches. The larger CSD source, prepared cache and deployment bundle receive distinct fingerprints.",
        69: "A local browser interface will accept an English MP4, MOV or WebM with one visible speaker, lasting at most 60 seconds and no larger than 200 MB. It will preview the video, show analysis progress, display the recognized transcript, positive/negative model probabilities and timestamped window predictions. Audio is converted to mono 16 kHz; local speech recognition, BERT, COVAREP and OpenFace supply the shared feature pipeline. The app rejects missing speech, insufficient usable face/voice coverage and incompatible checkpoints with clear messages. Per-request intermediate files are removed after processing, and uploaded preview files expire from the temporary cache.",
        79: "Implementation resources. Preferred CSD source: https://www.kaggle.com/datasets/samarwarsi/cmu-mosei. Official video folds: https://github.com/CMU-MultiComp-Lab/CMU-MultimodalSDK. Live extractors: https://github.com/covarep/covarep, https://github.com/TadasBaltrusaitis/OpenFace and https://github.com/SYSTRAN/faster-whisper. Colab runtime guidance: https://research.google.com/colaboratory/faq.html. The repository records pinned revisions, local hashes, schemas and actual execution evidence.",
    }
    for index, text in replacements.items():
        paragraph = doc.paragraphs[index]
        paragraph.text = text
    paragraph = doc.paragraphs[70].insert_paragraph_before(
        "Uploaded videos have no supplied ground-truth label. The demonstration therefore shows model outputs rather than correctness claims. Released word transcripts and automatically recognized speech differ, and the historical CSD extractor versions are not fully specified. Analyze those distribution differences and document observed recognition/tracking failures. Missing-modality experiments remain controlled benchmark diagnostics; the normal upload flow requires usable speech and face information. A small integration checkpoint may prove the pipeline works, but final performance claims require the complete frozen research evaluation."
    )
    paragraph.style = doc.styles["Normal"]
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.keep_together = True
    doc.tables[0].rows[3].cells[
        1
    ].text = "English CMU-MOSEI CSD streams with official video folds and shared word alignment"
    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        if style_name in doc.styles:
            doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    # Replace the existing architecture drawing in-place, keeping its page placement.
    canvas = Image.new("RGB", (1680, 820), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 35)
    small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 31)

    def box(bounds, text):
        draw.rounded_rectangle(bounds, radius=18, fill="#F0F4F6", outline="#526775", width=3)
        x0, y0, x1, y1 = bounds
        draw.multiline_text(
            ((x0 + x1) / 2, (y0 + y1) / 2),
            text,
            font=font,
            fill="black",
            anchor="mm",
            align="center",
            spacing=8,
        )

    def arrow(start, end):
        import math

        draw.line([start, end], fill="#526775", width=4)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        points = [end]
        for offset in (-0.45, 0.45):
            points.append(
                (end[0] - 20 * math.cos(angle + offset), end[1] - 20 * math.sin(angle + offset))
            )
        draw.polygon(points, fill="#526775")

    draw.text(
        (840, 35),
        "English video extraction or released CMU-MOSEI streams",
        font=small,
        fill="black",
        anchor="mm",
    )
    box((30, 100, 535, 260), "Text\nBERT word features\n768 dimensions")
    box((588, 100, 1092, 260), "Audio\nCOVAREP features\n73 dimensions")
    box((1145, 100, 1650, 260), "Vision\nOpenFace features\n40 dimensions")
    box((380, 355, 1300, 490), "Shared word alignment and masks\nWindows of at most 50 words")
    for x in (282, 840, 1398):
        arrow((x, 270), (max(430, min(x, 1250)), 345))
    arrow((840, 500), (840, 595))
    box((30, 605, 535, 790), "Modality projections\nand position\ninformation")
    box((588, 605, 1092, 790), "Transformer fusion\nWindow classification\nlogits")
    box((1145, 605, 1650, 790), "Duration weighted\nclip logits\nPositive or negative")
    # Route the aligned streams through projections before joint attention.
    draw.line([(840, 540), (282, 540), (282, 585)], fill="#526775", width=4)
    draw.rectangle((815, 545, 865, 603), fill="white")
    arrow((282, 575), (282, 595))
    arrow((545, 697), (578, 697))
    arrow((1102, 697), (1135, 697))
    image = io.BytesIO()
    canvas.save(image, format="PNG")
    shape = doc.inline_shapes[0]
    shape.height = int(shape.width * canvas.height / canvas.width)
    relation = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[relation]._blob = image.getvalue()
    doc.save(PROPOSAL)
    print(PROPOSAL)


if __name__ == "__main__":
    main()
